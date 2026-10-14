from datetime import timedelta
from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from hubflow.boards.models import Card
from hubflow.sprints.models import Sprint
from hubflow.timetracking.models import TimeEntry

pytestmark = pytest.mark.django_db


@pytest.fixture
def data(loja, user):
    sprint1 = Sprint.objects.create(project=loja)
    sprint2 = Sprint.objects.create(project=loja)
    a = Card.objects.create(project=loja, title='A', sprint=sprint1)
    b = Card.objects.create(project=loja, title='B', sprint=sprint2)
    d1 = timezone.make_aware(timezone.datetime(2026, 10, 1, 9, 0))
    d2 = timezone.make_aware(timezone.datetime(2026, 10, 2, 14, 0))
    d3 = timezone.make_aware(timezone.datetime(2026, 11, 3, 9, 0))
    TimeEntry.objects.create(card=a, user=user, started_at=d1, ended_at=d1 + timedelta(hours=2))
    TimeEntry.objects.create(card=a, user=user, started_at=d2, ended_at=d2 + timedelta(hours=1, minutes=30))
    TimeEntry.objects.create(card=b, user=user, started_at=d3, ended_at=d3 + timedelta(hours=1))
    TimeEntry.objects.create(card=b, user=user, started_at=d3 + timedelta(hours=5))  # aberta: fica de fora
    return {'a': a, 'b': b}


def export(api, loja, **params):
    query = '&'.join(f'{k}={v}' for k, v in params.items())
    return api.get(f'/api/projects/{loja.id}/timesheet.xlsx' + (f'?{query}' if query else ''))


def test_devolve_xlsx_com_as_abas(api, loja, data):
    response = export(api, loja)

    assert response.status_code == 200
    assert response['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'timesheet-loja' in response['Content-Disposition']
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ['Entradas', 'Por dia', 'Por mês', 'Por sprint']


def test_aba_de_entradas_lista_as_fechadas(api, loja, data):
    sheet = load_workbook(BytesIO(export(api, loja).content))['Entradas']
    rows = list(sheet.iter_rows(values_only=True))

    assert rows[0] == ('Data', 'Início', 'Fim', 'Duração (h)', 'Card', 'Título', 'Sprint', 'Usuário', 'Nota')
    assert len(rows) == 4  # cabeçalho + 3 fechadas
    assert rows[1][4] == '#1' and rows[1][3] == 2.0 and rows[1][6] == 1


def test_totais_por_dia_mes_e_sprint(api, loja, data):
    workbook = load_workbook(BytesIO(export(api, loja).content))

    por_dia = list(workbook['Por dia'].iter_rows(values_only=True))
    por_mes = list(workbook['Por mês'].iter_rows(values_only=True))
    por_sprint = list(workbook['Por sprint'].iter_rows(values_only=True))

    assert por_dia[1:] == [('2026-10-01', 2.0), ('2026-10-02', 1.5), ('2026-11-03', 1.0)]
    assert por_mes[1:] == [('2026-10', 3.5), ('2026-11', 1.0)]
    assert por_sprint[1:] == [('Sprint 1', 3.5), ('Sprint 2', 1.0)]


def test_filtra_por_periodo(api, loja, data):
    sheet = load_workbook(BytesIO(export(api, loja, **{'from': '2026-10-02', 'to': '2026-10-31'}).content))['Entradas']

    assert len(list(sheet.iter_rows(values_only=True))) == 2


def test_periodo_invalido_retorna_400(api, loja, data):
    assert export(api, loja, **{'from': '2026-10-31', 'to': '2026-10-01'}).status_code == 400


def test_projeto_inexistente_retorna_404(api):
    assert api.get('/api/projects/999/timesheet.xlsx').status_code == 404
